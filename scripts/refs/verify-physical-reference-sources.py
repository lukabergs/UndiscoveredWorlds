# /// script
# dependencies = ["numpy", "netcdf4", "rasterio", "pyhdf", "openpyxl"]
# ///
"""Read bounded samples of physical sources and check archive integrity, without extraction."""

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import tarfile
import zipfile

import netCDF4
import numpy as np
import openpyxl
from pyhdf.SD import SD, SDC
import rasterio

REPO = Path(__file__).resolve().parents[2]


def stats(values):
    array = np.ma.asarray(values)
    raw = np.asarray(array.data)
    valid = raw[(~np.ma.getmaskarray(array)) & np.isfinite(raw)]
    return dict(sample_valid=int(valid.size),
                sample_min=float(valid.min()) if valid.size else None,
                sample_max=float(valid.max()) if valid.size else None)


def inspect(path, kind):
    if kind == "tiff":
        with rasterio.Env(GDAL_CACHEMAX=64 * 1024**2), rasterio.open(path) as data:
            result = dict(shape=[data.height, data.width], crs=str(data.crs),
                          bounds=list(data.bounds), nodata=data.nodata,
                          **stats(data.read(1, out_shape=(180, 360), masked=True)))
            if not data.crs or not result['sample_valid']:
                raise ValueError("Raster has no CRS or no valid sampled pixels")
            if path.parent.name == 'soilgrids-2-wcs':
                if result['shape'] != [580, 1440] or not np.allclose(result['bounds'], [-180, -60, 180, 85]):
                    raise ValueError('SoilGrids WCS coverage differs from request')
            return result
    if kind == "netcdf":
        with netCDF4.Dataset(path) as data:
            result = dict(dimensions={k: len(v) for k, v in data.dimensions.items()}, variables={})
            for name, var in data.variables.items():
                if np.dtype(var.dtype).kind not in 'iuf':
                    continue
                key = tuple(slice(0, n, max(1, n // 64)) for n in var.shape)
                result['variables'][name] = dict(shape=list(var.shape), units=getattr(var, 'units', None),
                                                 **stats(var[key]))
            if not result['variables']:
                raise ValueError('No readable numeric NetCDF variables')
            return result
    if kind == "hdf":
        data = SD(str(path), SDC.READ)
        result = {}
        try:
            for name in ('Gpp_500m', 'Npp_500m', 'Npp_QC_500m'):
                var = data.select(name)
                try:
                    attrs = var.attributes()
                    values = var.get(start=(0, 0), count=(40, 40), stride=(60, 60))
                    if '_FillValue' in attrs:
                        values = np.ma.masked_equal(values, attrs['_FillValue'])
                    if 'valid_range' in attrs:
                        values = np.ma.masked_outside(values, *attrs['valid_range'])
                    result[name] = dict(shape=list(var.info()[2]), units=attrs.get('units'),
                                        fill=attrs.get('_FillValue'), scale=attrs.get('scale_factor'),
                                        valid_range=attrs.get('valid_range'), **stats(values))
                    if list(var.info()[2]) != [2400, 2400]:
                        raise ValueError('Unexpected MOD17 tile dimensions')
                finally:
                    var.endaccess()
        finally:
            data.end()
        return result
    if kind in ('zip', 'xlsx'):
        with zipfile.ZipFile(path) as data:
            unpacked = sum(x.file_size for x in data.infolist())
            if unpacked > 100 * 1024**3:
                raise ValueError('Archive exceeds bounded integrity-check size')
            bad = data.testzip()
            if bad:
                raise ValueError('ZIP CRC failure: ' + bad)
            result = dict(members=len(data.infolist()), uncompressed_bytes=unpacked, crc='passed')
            dbfs = [x for x in data.namelist() if x.lower().endswith('.dbf')]
            result['dbf_records'] = {}
            for name in dbfs:
                with data.open(name) as stream:
                    header = stream.read(32)
                result['dbf_records'][name] = int.from_bytes(header[4:8], 'little')
        if kind == 'xlsx':
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                result['sheets'] = {sheet.title: dict(rows=sheet.max_row, columns=sheet.max_column,
                    header_sample=list(sheet.iter_rows(min_row=1, max_row=3, max_col=12, values_only=True)))
                    for sheet in workbook}
            finally:
                workbook.close()
        return result
    if kind == 'tar':
        count = 0
        with tarfile.open(path, 'r|gz') as data:
            for member in data:
                count += 1
                if member.isfile():
                    with data.extractfile(member) as stream:
                        while stream.read(1024**2):
                            pass
        return dict(members=count, full_stream_read=True)
    if kind == 'geojson':
        data = json.loads(path.read_text(encoding='utf-8'))
        features = data['features']
        return dict(features=len(features), fields=list(features[0]['properties']),
                    geometries=sum(f.get('geometry') is not None for f in features))
    if path.name == 'ldem_4.img':
        label = path.with_suffix('.lbl').read_text()
        fields = {key: int(re.search(r'\b' + key + r'\s*=\s*(\d+)', label).group(1))
                  for key in ('LINES', 'LINE_SAMPLES', 'SAMPLE_BITS')}
        if path.stat().st_size != fields['LINES'] * fields['LINE_SAMPLES'] * fields['SAMPLE_BITS'] // 8:
            raise ValueError('LOLA image size differs from PDS label')
        return fields
    return dict(check='Format and SHA256 checked by downloader; no semantic reader for this format')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=REPO / 'refs/source/physical')
    parser.add_argument('--dataset', action='append')
    args = parser.parse_args()
    entries = json.loads((REPO / 'configs/physical-reference-sources.json').read_text())['files']
    results = []
    for entry in entries:
        if args.dataset and entry['dataset'] not in args.dataset:
            continue
        path = args.output / entry['dataset'] / entry['file']
        result = dict(dataset=entry['dataset'], file=entry['file'])
        try:
            receipt = json.loads(path.with_name(path.name + '.receipt.json').read_text())
            with path.open('rb') as stream:
                checksum = hashlib.file_digest(stream, 'sha256').hexdigest()
            if receipt['sha256'] != checksum or receipt['request'] != entry:
                raise ValueError('Receipt does not match source or catalog')
            result.update(status='passed', sha256=checksum, details=inspect(path, entry['kind']))
        except Exception as exc:
            result.update(status='failed', error=str(exc))
        results.append(result)
        print(f"{result['status']}: {entry['dataset']}/{entry['file']}", flush=True)
    report = dict(checked_utc=datetime.now(timezone.utc).isoformat(),
                  scope='Archive integrity and bounded data samples; not scientific accuracy or simulator validation.',
                  results=results)
    (args.output / 'source-verification.json').write_text(json.dumps(report, indent=2, default=str) + '\n')
    return int(any(r['status'] != 'passed' for r in results))


if __name__ == '__main__':
    raise SystemExit(main())
